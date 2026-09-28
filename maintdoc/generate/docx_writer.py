"""DOCX manual writer (python-docx) - same structured content as the PDF."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

from maintdoc.constants import TOOL_VERSION
from maintdoc.generate.manual import Item, Manual

log = logging.getLogger(__name__)

ADMONITION_FILL = {"danger": "F8D7DA", "warning": "FFE5CC", "caution": "FFF3CD", "notice": "DBE9FF", "note": "EEEEEE"}
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB0, 0x00, 0x20)


def _clean(text: str | None) -> str:
    # XML 1.0 forbids most control characters
    return "".join(ch for ch in (text or "") if ch in "\t\n\r" or ord(ch) >= 32)


def _field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cite_run(p, item: Item) -> None:
    r = p.add_run(" " + item.citation_text())
    r.font.size = Pt(7)
    r.font.color.rgb = GREY
    notes = list(item.flags) + [f"CONFLICT {c}" for c in item.conflicts]
    if notes:
        r2 = p.add_run(" {" + "; ".join(notes) + "}")
        r2.font.size = Pt(7)
        r2.font.color.rgb = RED


def _table(doc, rows: list[list], header: bool = True, size: float = 8.0):
    """Build the table row by row at XML level.

    python-docx's ``table.cell(i, j)`` recomputes the whole cell grid on every call, which is
    quadratic for large tables (e.g. a citation index with thousands of rows).
    """
    ncols = max((len(r) for r in rows), default=1)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        tr = t._tbl.add_tr()
        for j in range(ncols):
            val = r[j] if j < len(r) else ""
            cell = _Cell(tr.add_tc(), t)
            para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
            run = para.add_run(_clean("" if val is None else str(val)))
            run.font.size = Pt(size)
            if header and i == 0:
                run.bold = True
                _shade(cell, "E8E8E8")
    return t


def _item(doc, item: Item, cfg) -> None:
    if item.kind == "admonition":
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.cell(0, 0)
        _shade(cell, ADMONITION_FILL.get(item.safety_level or item.block_type, "EEEEEE"))
        p = cell.paragraphs[0]
        run = p.add_run(_clean(item.text))
        run.bold = item.block_type in ("danger", "warning")
        _cite_run(p, item)
        doc.add_paragraph()
    elif item.kind == "table" and item.table_rows:
        if item.caption:
            cp = doc.add_paragraph()
            cr = cp.add_run(_clean(item.caption))
            cr.bold = True
        _table(doc, item.table_rows)
        p = doc.add_paragraph()
        _cite_run(p, item)
    elif item.kind == "figure":
        if item.figure_path and Path(item.figure_path).exists():
            try:
                doc.add_picture(item.figure_path, width=Cm(float(cfg.get("generation.figure_max_width_cm", 15.0))))
            except Exception as exc:  # noqa: BLE001
                doc.add_paragraph(f"[Figure image could not be embedded: {exc}]")
        else:
            doc.add_paragraph("[Figure image not available - see source page]")
        p = doc.add_paragraph()
        p.add_run(_clean(item.caption or "(figure without caption)")).italic = True
        _cite_run(p, item)
        note = doc.add_paragraph("Figure reproduced from the source page; human visual check required.")
        note.runs[0].font.size = Pt(7)
    else:
        p = doc.add_paragraph(style="List Paragraph" if item.kind == "list_item" else None)
        p.add_run(_clean(item.text))
        _cite_run(p, item)


def write_docx(manual: Manual, path: Path, cfg) -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    banner = "DRAFT - NOT RELEASED" if manual.mode == "draft" else \
        "APPROVED CONTENT - NOT RELEASED UNTIL FORMAL ENGINEERING RELEASE"
    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{banner}    {manual.title}")
    hr.font.size = Pt(7.5)
    hr.font.color.rgb = RED
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"Generated {manual.generated_at} by maintdoc {TOOL_VERSION} - page ")
    fr.font.size = Pt(7.5)
    _field(fp, "PAGE")
    doc.core_properties.title = f"{manual.title} ({manual.mode.upper()})"
    doc.core_properties.author = "maintdoc (rule-based, not released)"
    doc.core_properties.comments = manual.disclaimer[:250]  # OOXML core property limit is 255 chars

    doc.add_heading(_clean(manual.title), level=0)
    p = doc.add_paragraph()
    r = p.add_run(banner)
    r.bold = True
    r.font.color.rgb = RED
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if manual.organisation:
        doc.add_paragraph(_clean(manual.organisation))
    doc.add_paragraph(f"Generated: {manual.generated_at} - maintdoc {TOOL_VERSION} - run {manual.run_id} - "
                      f"configuration {manual.config_hash}")
    d = doc.add_paragraph()
    dr = d.add_run(manual.disclaimer)
    dr.bold = True
    dr.font.color.rgb = RED
    _table(doc, [["Measure", "Value"]] + [[k.replace("_", " "), str(v)] for k, v in manual.stats.items()])
    doc.add_page_break()
    doc.add_paragraph("Contents (right-click and choose 'Update Field' in Word)").runs[0].bold = True
    _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    doc.add_page_break()
    doc.add_heading("Document control", level=1)
    doc.add_paragraph("Each statement is reproduced verbatim from its source (or in reviewer wording that preserves "
                      "every number and unit) and is followed by its evidence citation. The citation index maps each "
                      "citation to file, page, section, SHA-256 and extraction method. Figures and complex layouts "
                      "require a human visual check against the source pages.")
    for n, ch in enumerate(manual.chapters, start=1):
        doc.add_page_break()
        doc.add_heading(f"{n}. {ch.title}", level=1)
        if not ch.topics:
            doc.add_paragraph("No approved content for this chapter." if manual.mode == "approved"
                              else "No content was classified into this chapter.")
        for t in ch.topics:
            doc.add_heading(_clean(t.heading), level=2)
            cp = doc.add_paragraph()
            cr = cp.add_run(f"Applicability: {t.applicability} - section heading evidence: "
                            f"{', '.join(t.heading_citations) or 'none'}")
            cr.font.size = Pt(7)
            cr.font.color.rgb = GREY
            multi = len({i.source_label for i in t.items}) > 1
            last = None
            for item in t.items:
                if multi and item.source_label != last:
                    sp = doc.add_paragraph()
                    sr = sp.add_run(f"Source: {item.source_label}")
                    sr.font.size = Pt(7)
                    sr.italic = True
                    last = item.source_label
                _item(doc, item, cfg)
        for gt in ch.generated_tables:
            doc.add_heading(_clean(gt["title"]), level=3)
            cols = gt["columns"]
            _table(doc, [cols] + [[r.get(c, "") for c in cols] for r in gt["rows"]], size=7)
    doc.add_page_break()
    doc.add_heading("Appendix A - Conflicts", level=1)
    if manual.conflicts:
        _table(doc, [["Conflict", "Type / severity", "Status", "Description", "Resolution"]] +
               [[c["conflict_id"], f"{c['conflict_type']} / {c['severity']}", c["status"], c["description"],
                 f"{c.get('resolution_type') or ''} {c.get('resolution') or ''} {c.get('resolved_by') or ''}"]
                for c in manual.conflicts], size=7)
    else:
        doc.add_paragraph("None.")
    doc.add_heading("Appendix B - Withheld items", level=1)
    if manual.withheld:
        _table(doc, [["Evidence", "Reason"]] + [[w["evidence_id"], w["reason"]] for w in manual.withheld], size=7)
    else:
        doc.add_paragraph("None.")
    doc.add_heading("Appendix C - Citation index", level=1)
    rows = [["Evidence ID", "Source / file", "Page", "Section", "Method / OCR", "Review", "SHA-256"]]
    for eid in sorted(manual.citation_index):
        c = manual.citation_index[eid]
        rows.append([eid, f"{c.source_id} {c.filename}" + (f" Rev {c.revision}" if c.revision else ""), str(c.page_no),
                     (c.section or "")[:120], c.method + (f" {c.ocr_confidence}%" if c.ocr_confidence else ""),
                     c.review_status, c.sha256])
    _table(doc, rows, size=6.5)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.docx")
    doc.save(str(tmp))
    tmp.replace(path)
    return path
