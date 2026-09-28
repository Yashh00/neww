"""Verification of source PDFs, extracted evidence and generated manuals.

Produces one result row per check (pass / fail / warn / pending_human /
skipped) plus the acceptance-criteria summary. Verification never modifies
review state; it only reports. Passing verification does NOT mean the
content is released: formal engineering release is always a separate,
human step.
"""

from __future__ import annotations

import json
import logging
import re
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

from maintdoc import audit
from maintdoc.analysis.units import UnitExtractor, quantity_signature
from maintdoc.approval import approval_blockers, effective_text, fingerprint, numeric_identity
from maintdoc.config import Config
from maintdoc.constants import (ExtractionStatus, PageStatus, ReviewStatus, SourceStatus, TOOL_VERSION,
                                VisualCheckStatus)
from maintdoc.db import rows, transaction
from maintdoc.generate.run import output_paths
from maintdoc.utils import loads, now_iso, numeric_tokens, sha256_file

log = logging.getLogger(__name__)

PASS, FAIL, WARN, PENDING, SKIP = "pass", "fail", "warn", "pending_human", "skipped"
_WS = re.compile(r"\s+")


def squash(text: str) -> str:
    """Whitespace-free, NFKC text for layout-independent containment checks."""
    import unicodedata
    return _WS.sub("", unicodedata.normalize("NFKC", text or ""))


@dataclass
class Check:
    group: str
    name: str
    target: str
    status: str
    detail: str = ""


@dataclass
class VerifyReport:
    run_id: str
    checks: list[Check] = field(default_factory=list)
    acceptance: list[dict[str, str]] = field(default_factory=list)
    overall: str = ""
    generated_at: str = field(default_factory=now_iso)

    def add(self, group: str, name: str, target: str, status: str, detail: str = "") -> None:
        self.checks.append(Check(group, name, target, status, detail))

    def counts(self) -> dict[str, int]:
        return dict(Counter(c.status for c in self.checks))

    def group_status(self, group: str, name: str | None = None) -> str:
        sts = [c.status for c in self.checks if c.group == group and (name is None or c.name == name)]
        for s in (FAIL, PENDING, WARN):
            if s in sts:
                return s
        return PASS if sts else SKIP


# --------------------------------------------------------------------------- sources & pages
def check_sources(conn: sqlite3.Connection, cfg: Config, rep: VerifyReport, rehash: bool) -> None:
    use_plumber = bool(cfg.get("verify.pdfplumber_page_count", True))
    for s in rows(conn, "SELECT * FROM sources ORDER BY source_id"):
        sid = s["source_id"]
        final = s["status"] != SourceStatus.OK or s["extraction_status"] in ExtractionStatus.FINAL
        rep.add("sources", "processing_status", sid, PASS if (s["status"] and final) else FAIL,
                f"{s['status']}/{s['extraction_status']}" + ("" if final else " - not yet processed (run extract)"))
        if not s["present"]:
            rep.add("sources", "file_present", sid, WARN, "registered file missing (approvals invalidated)")
            continue
        path = Path(s["abs_path"])
        if s["status"] in (SourceStatus.OFFLINE_PLACEHOLDER, SourceStatus.UNREADABLE):
            rep.add("sources", "file_readable", sid, FAIL, f"{s['status']}: {s['status_detail']}")
            continue
        if rehash and s["sha256"]:
            try:
                ok = sha256_file(path) == s["sha256"]
                rep.add("sources", "sha256_unchanged", sid, PASS if ok else FAIL,
                        "" if ok else "file changed since inventory - run inventory (approvals will be invalidated)")
            except OSError as exc:
                rep.add("sources", "sha256_unchanged", sid, FAIL, str(exc))
                continue
        if s["status"] not in (SourceStatus.OK, SourceStatus.DUPLICATE):
            rep.add("sources", "readable_page_count", sid, SKIP, f"status {s['status']}")
            continue
        try:
            with pymupdf.open(str(path)) as d:
                n = d.page_count
            detail = f"PyMuPDF {n}, register {s['page_count']}"
            status = PASS if n == s["page_count"] else FAIL
            if use_plumber:
                try:
                    import pdfplumber
                    with pdfplumber.open(str(path)) as pdf:
                        n2 = len(pdf.pages)
                    detail += f", pdfplumber {n2}"
                    if n2 != n:
                        status = WARN if status == PASS else status
                except Exception as exc:  # noqa: BLE001
                    detail += f", pdfplumber failed: {exc}"
                    status = WARN if status == PASS else status
            rep.add("sources", "readable_page_count", sid, status, detail)
        except Exception as exc:  # noqa: BLE001
            rep.add("sources", "readable_page_count", sid, FAIL, f"cannot open: {exc}")


def check_pages(conn: sqlite3.Connection, rep: VerifyReport) -> None:
    for s in rows(conn, "SELECT * FROM sources WHERE status=? AND present=1 ORDER BY source_id", (SourceStatus.OK,)):
        sid = s["source_id"]
        if s["extraction_status"] not in ExtractionStatus.FINAL:
            rep.add("pages", "every_page_has_status", sid, FAIL, "source not extracted")
            continue
        pages = rows(conn, "SELECT page_no, extraction_status, source_sha256 FROM pages WHERE source_id=? "
                           "ORDER BY page_no", (sid,))
        nums = [p["page_no"] for p in pages]
        expected = list(range(1, (s["page_count"] or 0) + 1))
        missing = sorted(set(expected) - set(nums))
        no_status = [p["page_no"] for p in pages if not p["extraction_status"]]
        stale = [p["page_no"] for p in pages if p["source_sha256"] != s["sha256"]]
        ok = not missing and not no_status and not stale
        rep.add("pages", "every_page_has_status", sid, PASS if ok else FAIL,
                f"{len(pages)}/{len(expected)} pages" + (f"; missing {missing[:20]}" if missing else "") +
                (f"; stale {stale[:20]}" if stale else ""))
        st = Counter(p["extraction_status"] for p in pages)
        hard = {k: v for k, v in st.items() if k in (PageStatus.FAILED, PageStatus.OCR_UNAVAILABLE, PageStatus.OCR_FAILED,
                                                    PageStatus.LOW_QUALITY_TEXT)}
        soft = {k: v for k, v in st.items() if k in (PageStatus.EMPTY, PageStatus.GRAPHIC_ONLY,
                                                    PageStatus.OCR_LOW_CONFIDENCE)}
        rep.add("pages", "extracted_page_completeness", sid, FAIL if hard else (WARN if soft else PASS),
                ", ".join(f"{k}={v}" for k, v in sorted(st.items())))


def check_evidence_text(conn: sqlite3.Connection, cfg: Config, rep: VerifyReport) -> None:
    scope = cfg.get("verify.recheck_evidence_text", "approved")
    if scope == "none":
        rep.add("evidence", "source_text_recheck", "all", SKIP, "disabled in configuration")
        return
    where = "e.status='active' AND e.in_manual=1" + (" AND e.review_status='approved'" if scope == "approved" else "")
    items = rows(conn, f"SELECT e.*, s.abs_path, s.sha256 cur_sha FROM evidence e JOIN sources s ON "
                       f"s.source_id=e.source_id WHERE {where} ORDER BY e.source_id, e.page_no")
    by_page: dict[tuple[str, int], list[dict]] = defaultdict(list)
    for e in items:
        by_page[(e["abs_path"], e["page_no"])].append(e)
    checked = failed = skipped = 0
    for (path, pno), evs in by_page.items():
        try:
            with pymupdf.open(path) as d:
                page_text = squash(d.load_page(pno - 1).get_text("text"))
        except Exception as exc:  # noqa: BLE001
            for e in evs:
                rep.add("evidence", "source_text_recheck", e["evidence_id"], FAIL, f"cannot read page: {exc}")
                failed += 1
            continue
        for e in evs:
            if e["extraction_method"] in ("tesseract_ocr", "pymupdf_image", "pymupdf_drawing"):
                skipped += 1
                continue  # OCR/figure evidence is confirmed by humans (ocr_verified / visual checks)
            if e["source_sha256"] != e["cur_sha"]:
                rep.add("evidence", "source_text_recheck", e["evidence_id"], FAIL, "source changed")
                failed += 1
                continue
            parts = [c for r in (loads(e["table_json"], []) or []) for c in r if c] if e["table_json"] else [e["text"]]
            missing = [p for p in parts if squash(p) not in page_text]
            checked += 1
            if missing:
                failed += 1
                rep.add("evidence", "source_text_recheck", e["evidence_id"], FAIL,
                        f"text not found on source page {pno}: {missing[0][:80]!r}")
    rep.add("evidence", "source_text_recheck_summary", scope, FAIL if failed else PASS,
            f"{checked} re-checked against source pages, {failed} failed, {skipped} OCR/figure items rely on human checks")


# --------------------------------------------------------------------------- approvals
def check_approvals(conn: sqlite3.Connection, cfg: Config, rep: VerifyReport) -> dict[str, int]:
    ux = UnitExtractor(cfg)
    counts = Counter()
    for e in rows(conn, "SELECT * FROM evidence WHERE review_status=?", (ReviewStatus.APPROVED,)):
        eid = e["evidence_id"]
        counts["approved"] += 1
        blockers = approval_blockers(conn, cfg, eid)
        fp_ok = fingerprint(conn, e) == e["approval_fingerprint"]
        cit_problems = [b for b in blockers if any(k in b for k in ("source", "cited evidence", "evidence is",
                                                                     "does not exist"))]
        rep.add("approvals", "citation_integrity", eid, FAIL if cit_problems else PASS, "; ".join(cit_problems))
        if cit_problems:
            counts["citation_fail"] += 1
        num = numeric_identity(e["text"], e.get("display_text"), cfg)
        stored = rows(conn, "SELECT value_text, unit_norm FROM quantities WHERE evidence_id=?", (eid,))
        fresh = quantity_signature(ux.extract(e["text"], e["section_heading"],
                                              loads(e["table_json"]) if e["table_json"] else None))
        stored_sig = sorted(((r["value_text"] or "").replace("−", "-"), r["unit_norm"]) for r in stored)
        if stored_sig != fresh:
            num.append("stored values differ from re-extraction (run validate)")
        rep.add("approvals", "numeric_identity", eid, FAIL if num else PASS, "; ".join(num))
        if num:
            counts["numeric_fail"] += 1
        conflict = [b for b in blockers if "conflict" in b]
        rep.add("approvals", "no_open_blocking_conflict", eid, FAIL if conflict else PASS, "; ".join(conflict))
        if conflict:
            counts["conflict_fail"] += 1
        rep.add("approvals", "approval_still_valid", eid, PASS if (fp_ok and not blockers) else FAIL,
                "" if fp_ok else "fingerprint changed (source/evidence/wording changed) - run validate")
        if not (fp_ok and not blockers):
            counts["invalid"] += 1
    if not counts["approved"]:
        rep.add("approvals", "approved_statements", "all", WARN, "no statements approved yet")
    return dict(counts)


# --------------------------------------------------------------------------- generated manuals
def _body_text(page: pymupdf.Page) -> str:
    """Page text without the running header/footer, so statements that continue across a page
    break are contiguous in the extracted text (header/footer lie outside the body frame)."""
    margin = 15 * 72 / 25.4  # 15 mm; the body frame starts 22 mm from the top and ends 20 mm above the bottom
    return page.get_text("text", clip=pymupdf.Rect(0, margin, page.rect.width, page.rect.height - margin))


def _docx_text(path: Path) -> str:
    from docx import Document
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def check_outputs(conn: sqlite3.Connection, cfg: Config, rep: VerifyReport) -> None:
    open_blocking = {r[0] for r in conn.execute("SELECT conflict_id FROM conflicts WHERE status='open' AND blocking=1")}
    for mode in cfg.get("generation.modes", ["draft", "approved"]):
        paths = output_paths(cfg, mode)
        if not paths["manifest"].exists():
            last = conn.execute("SELECT status, detail FROM generated_outputs WHERE mode=? ORDER BY output_id DESC "
                                "LIMIT 1", (mode,)).fetchone()
            if last and last["status"] == "blocked":
                rep.add("generation", "manual_generated", mode, WARN if mode == "approved" else FAIL,
                        f"blocked: {last['detail'][:300]}")
            else:
                rep.add("generation", "manual_generated", mode, WARN if mode == "approved" else FAIL,
                        "not generated (run generate)")
            continue
        man = json.loads(paths["manifest"].read_text(encoding="utf-8"))
        items = man.get("items", [])
        texts: dict[str, str] = {}
        for kind in ("pdf", "docx"):
            p = paths[kind]
            if not p.exists():
                rep.add("generation", f"{kind}_readable", mode, FAIL, "file missing")
                continue
            try:
                if kind == "pdf":
                    with pymupdf.open(str(p)) as d:
                        n = d.page_count
                        texts[kind] = "\n".join(_body_text(pg) for pg in d)
                    rep.add("generation", "pdf_readable", mode, PASS if n > 0 else FAIL, f"{n} pages")
                else:
                    texts[kind] = _docx_text(p)
                    rep.add("generation", "docx_readable", mode, PASS, f"{len(texts[kind])} characters")
            except Exception as exc:  # noqa: BLE001
                rep.add("generation", f"{kind}_readable", mode, FAIL, str(exc))
        gen_errors = conn.execute("SELECT COUNT(*) FROM extraction_errors WHERE stage='generation' AND status='open'"
                                  ).fetchone()[0]
        rep.add("generation", "generation_errors", mode, FAIL if gen_errors else PASS, f"{gen_errors} open")
        # citations and numeric identity of every rendered item
        ev_ids = {r[0]: r for r in conn.execute("SELECT evidence_id, status, source_sha256, source_id, text, "
                                                "display_text FROM evidence")}
        cur_sha = {r[0]: r[1] for r in conn.execute("SELECT source_id, sha256 FROM sources")}
        uncited = bad_cite = text_missing = num_mismatch = 0
        squashed = {k: squash(v) for k, v in texts.items()}          # computed once per output file
        output_numbers = {k: numeric_tokens(v) for k, v in texts.items()}
        for it in items:
            if not it.get("citations"):
                uncited += 1
                continue
            for cid in it["citations"]:
                r = ev_ids.get(cid)
                if r is None or r["status"] != "active" or r["source_sha256"] != cur_sha.get(r["source_id"]):
                    bad_cite += 1
                    rep.add("generation", "citation_integrity", f"{mode}:{cid}", FAIL,
                            "cited evidence missing, superseded or source changed")
            r = ev_ids.get(it["item_id"])
            if r is not None and it["kind"] not in ("table", "figure"):
                if it["text"] != effective_text(dict(r)):
                    num_mismatch += 1
                    rep.add("generation", "manifest_matches_evidence", f"{mode}:{it['item_id']}", FAIL,
                            "rendered text differs from approved/original evidence text")
            for kind, full in texts.items():
                sq = squashed[kind]
                needles = [it["citation_text"]]
                if it["kind"] == "table" and it.get("table_rows"):
                    needles += [c for row_ in it["table_rows"] for c in row_ if c]
                elif it["text"]:
                    needles.append(it["text"])
                missing = [n for n in needles if squash(n) not in sq]
                if missing:
                    text_missing += 1
                    rep.add("generation", f"{kind}_content_present", f"{mode}:{it['item_id']}", FAIL,
                            f"not found in {kind}: {missing[0][:80]!r}")
                elif Counter(it["numbers"]) - output_numbers[kind]:
                    num_mismatch += 1
                    rep.add("generation", f"{kind}_numeric_identity", f"{mode}:{it['item_id']}", FAIL,
                            "numbers of the item not all present in output")
        rep.add("generation", "every_item_cited", mode, FAIL if uncited else PASS,
                f"{len(items)} items, {uncited} without citation")
        rep.add("generation", "citations_valid", mode, FAIL if bad_cite else PASS, f"{bad_cite} invalid citations")
        rep.add("generation", "content_and_numbers_present", mode, FAIL if (text_missing or num_mismatch) else PASS,
                f"{text_missing} items missing text, {num_mismatch} numeric/identity mismatches")
        # expected sections
        full_pdf = texts.get("pdf", "")
        chapters = {c["id"]: c for c in man.get("chapters", [])}
        for cid in cfg.get("verify.expected_chapters", []):
            c = chapters.get(cid)
            title = cfg.chapter_title(cid)
            if c is None or squash(title) not in squash(full_pdf):
                rep.add("generation", "expected_section", f"{mode}:{cid}", FAIL, f"section '{title}' missing")
            elif c["items"] == 0:
                rep.add("generation", "expected_section", f"{mode}:{cid}", WARN, f"section '{title}' has no content")
            else:
                rep.add("generation", "expected_section", f"{mode}:{cid}", PASS, f"{c['items']} items")
        # tables present
        n_tables = sum(1 for it in items if it["kind"] == "table")
        rep.add("generation", "tables_present", mode, PASS if n_tables or mode == "approved" else WARN,
                f"{n_tables} tables rendered")
        if mode == "approved":
            in_conflict = [it["item_id"] for it in items if set(it.get("conflicts", [])) & open_blocking]
            rep.add("generation", "approved_manual_free_of_blocking_conflicts", mode,
                    FAIL if (in_conflict or open_blocking) else PASS,
                    f"{len(open_blocking)} open blocking conflicts; {len(in_conflict)} rendered items affected"
                    + (" - approved manual is outdated; re-run generate" if open_blocking else ""))
        if mode == "draft":
            eligible = {r[0] for r in conn.execute(
                "SELECT evidence_id FROM evidence WHERE status='active' AND in_manual=1 AND (dup_role IS NULL OR "
                "dup_role='canonical') AND review_status != 'rejected'")}
            rendered = {it["item_id"] for it in items} | {w["evidence_id"] for w in man.get("withheld", [])}
            missing = sorted(eligible - rendered)
            rep.add("generation", "missing_content", mode, FAIL if missing else PASS,
                    f"{len(missing)} eligible statements not in draft manual (re-run generate)" +
                    (f": {missing[:10]}" if missing else ""))


def check_register_consistency(conn: sqlite3.Connection, rep: VerifyReport) -> None:
    unexplained = conn.execute("SELECT COUNT(*) FROM evidence WHERE status='active' AND in_manual=0 AND "
                               "(manual_exclusion IS NULL OR manual_exclusion='')").fetchone()[0]
    rep.add("evidence", "exclusions_explained", "all", FAIL if unexplained else PASS,
            f"{unexplained} active evidence items excluded without a recorded reason (run validate)")
    pending_q = conn.execute("SELECT COUNT(*) FROM evidence WHERE status='active' AND quantities_done=0 AND block_type "
                             "IN ('paragraph','list_item','table','danger','warning','caution','notice','note')"
                             ).fetchone()[0]
    rep.add("evidence", "values_extracted", "all", FAIL if pending_q else PASS,
            f"{pending_q} statements without value extraction (run validate)")


def check_visual(conn: sqlite3.Connection, rep: VerifyReport) -> dict[str, int]:
    st = {r[0]: r[1] for r in conn.execute("SELECT status, COUNT(*) FROM visual_checks WHERE status!='invalidated' "
                                           "GROUP BY status")}
    for r in rows(conn, "SELECT * FROM visual_checks WHERE status IN ('pending','failed') ORDER BY check_id"):
        target = r["output_path"] or f"{r['source_id']} p.{r['page_no']}"
        rep.add("visual", r["kind"], target, PENDING if r["status"] == VisualCheckStatus.PENDING else FAIL,
                (r["reason"] or "") + (f" - {r['comment']}" if r["comment"] else ""))
    if not st:
        rep.add("visual", "visual_checks", "all", PASS, "no visual checks required")
    return st


def check_audit(conn: sqlite3.Connection, rep: VerifyReport) -> None:
    ok, bad, n = audit.verify_chain(conn)
    rep.add("audit", "hash_chain", "audit_log", PASS if ok else FAIL,
            f"{n} entries" + ("" if ok else f"; first invalid entry {bad}"))
    orphan = conn.execute("SELECT COUNT(*) FROM review_decisions d LEFT JOIN audit_log a ON a.audit_id=d.audit_id "
                          "WHERE a.audit_id IS NULL").fetchone()[0]
    rep.add("audit", "decisions_logged", "review_decisions", FAIL if orphan else PASS,
            f"{orphan} review decisions without audit entry")
    bad_err = conn.execute("SELECT COUNT(*) FROM extraction_errors WHERE error_code IS NULL OR stage IS NULL").fetchone()[0]
    rep.add("audit", "errors_logged", "extraction_errors", FAIL if bad_err else PASS,
            f"{conn.execute('SELECT COUNT(*) FROM extraction_errors').fetchone()[0]} errors registered")
    for t in ("audit_log", "review_decisions"):
        trig = conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='trigger' AND tbl_name=?", (t,)).fetchone()[0]
        rep.add("audit", "append_only_enforced", t, PASS if trig >= 2 else FAIL, f"{trig} protective triggers")


# --------------------------------------------------------------------------- acceptance
def acceptance(conn: sqlite3.Connection, rep: VerifyReport, approvals: dict[str, int], visual: dict[str, int]) -> None:
    def a(no: int, text: str, status: str, detail: str) -> None:
        rep.acceptance.append({"no": str(no), "criterion": text, "status": status, "detail": detail})

    g = rep.group_status
    a(1, "Every discovered PDF has a processing status", g("sources", "processing_status"),
      f"{sum(1 for c in rep.checks if c.name == 'processing_status')} sources checked")
    a(2, "Every page has an extraction status", g("pages", "every_page_has_status"),
      "page register complete for all extracted sources" if g("pages", "every_page_has_status") == PASS
      else "see Pages sheet")
    n_app = approvals.get("approved", 0)
    a(3, "All approved claims have source citations",
      FAIL if approvals.get("citation_fail") else PASS, f"{n_app} approved; {approvals.get('citation_fail', 0)} failing")
    a(4, "All approved numbers are checked against evidence",
      FAIL if approvals.get("numeric_fail") else PASS, f"{n_app} approved; {approvals.get('numeric_fail', 0)} failing")
    blocked_ok = not approvals.get("conflict_fail") and g("generation", "approved_manual_free_of_blocking_conflicts") \
        != FAIL
    open_blocking = conn.execute("SELECT COUNT(*) FROM conflicts WHERE status='open' AND blocking=1").fetchone()[0]
    a(5, "Unresolved critical conflicts block approval", PASS if blocked_ok else FAIL,
      f"{open_blocking} unresolved blocking conflicts; approved statements affected: {approvals.get('conflict_fail', 0)}")
    changed = [c for c in rep.checks if c.name == "sha256_unchanged" and c.status == FAIL]
    a(6, "Source changes invalidate dependent approvals",
      FAIL if (changed or approvals.get("invalid")) else PASS,
      f"{len(changed)} files changed since inventory; {approvals.get('invalid', 0)} approvals no longer valid")
    a(7, "All errors and review decisions are logged", g("audit"), "audit hash chain and decision linkage")
    pend = visual.get("pending", 0)
    failed = visual.get("failed", 0)
    a(8, "Mandatory human visual checks completed", FAIL if failed else (PENDING if pend else PASS),
      f"{pend} pending, {failed} failed")
    unreviewed = conn.execute("SELECT COUNT(*) FROM evidence WHERE status='active' AND in_manual=1 AND "
                              "(dup_role IS NULL OR dup_role='canonical') AND review_status != 'approved' AND "
                              "review_status != 'rejected'").fetchone()[0]
    a(9, "All manual statements reviewed", PENDING if unreviewed else PASS, f"{unreviewed} statements awaiting review")
    a(10, "Formal engineering release", PENDING,
      "Outputs are drafts. Rule-based automation cannot establish semantic completeness or safety certification; "
      "release requires qualified engineering review and sign-off outside this tool.")
    sts = [x["status"] for x in rep.acceptance[:9]]
    if FAIL in sts or any(c.status == FAIL for c in rep.checks):
        rep.overall = "FAIL"
    elif PENDING in sts:
        rep.overall = "INCOMPLETE - HUMAN REVIEW PENDING"
    else:
        rep.overall = "PASS (content verified; NOT RELEASED)"


def run_verify(conn: sqlite3.Connection, cfg: Config, run_id: str, rehash: bool | None = None,
               persist: bool = True) -> VerifyReport:
    rep = VerifyReport(run_id)
    rehash = bool(cfg.get("verify.rehash_sources", True)) if rehash is None else rehash
    log.info("verify: sources")
    check_sources(conn, cfg, rep, rehash)
    log.info("verify: pages")
    check_pages(conn, rep)
    log.info("verify: evidence")
    check_evidence_text(conn, cfg, rep)
    check_register_consistency(conn, rep)
    log.info("verify: approvals")
    approvals = check_approvals(conn, cfg, rep)
    log.info("verify: generated outputs")
    check_outputs(conn, cfg, rep)
    visual = check_visual(conn, rep)
    check_audit(conn, rep)
    acceptance(conn, rep, approvals, visual)
    if persist:
        with transaction(conn):
            conn.executemany("INSERT INTO verification_results(run_id, check_group, check_name, target, status, detail, "
                             "created_at) VALUES(?,?,?,?,?,?,?)",
                             [(run_id, c.group, c.name, c.target, c.status, c.detail[:2000], rep.generated_at)
                              for c in rep.checks] +
                             [(run_id, "acceptance", f"criterion_{x['no']}", x["criterion"], x["status"], x["detail"],
                               rep.generated_at) for x in rep.acceptance])
            audit.append(conn, "system", "verify.completed", "verification", run_id,
                         after={"overall": rep.overall, "counts": rep.counts(), "tool_version": TOOL_VERSION},
                         run_id=run_id)
    return rep
