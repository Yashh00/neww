"""Excel exports (openpyxl write-only mode for large registers).

Cell values are written as plain text (never formulas) and illegal XML
characters are removed. Text longer than the Excel cell limit is truncated
with a pointer to the database, which always holds the full text.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from maintdoc.config import Config
from maintdoc.constants import DRAFT_DISCLAIMER, ERROR_CODES, TOOL_VERSION
from maintdoc.db import rows
from maintdoc.utils import now_iso, truncate

MAX_ROWS = 1_048_000
HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
STATUS_FILL = {"pass": "C6EFCE", "fail": "FFC7CE", "warn": "FFEB9C", "pending_human": "BDD7EE",
               "critical": "FFC7CE", "major": "FFEB9C", "error": "FFC7CE"}


class Sheet:
    def __init__(self, title: str, columns: list[str], data: Iterable[Any], widths: dict[str, int] | None = None,
                 status_col: str | None = None):
        self.title = title
        self.columns = columns
        self.data = data
        self.widths = widths or {}
        self.status_col = status_col


def _cell(ws, value: Any, max_chars: int, bold: bool = False, fill: str | None = None):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", truncate(value, max_chars))
    elif value is not None and not isinstance(value, (int, float)):
        value = str(value)
    c = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        c.data_type = "s"  # never interpret text as a formula
    if bold:
        c.font = Font(bold=True)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def write_workbook(path: Path, sheets: list[Sheet], max_chars: int = 32000) -> Path:
    wb = Workbook(write_only=True)
    for sh in sheets:
        part = 1
        ws = None
        n = 0
        status_idx = sh.columns.index(sh.status_col) if sh.status_col in sh.columns else None

        def new_ws(p: int):
            title = sh.title[:28] + (f"_{p}" if p > 1 else "")
            w = wb.create_sheet(title=title)
            for i, col in enumerate(sh.columns, start=1):
                w.column_dimensions[get_column_letter(i)].width = sh.widths.get(col, min(60, max(10, len(col) + 4)))
            w.freeze_panes = "A2"
            w.append([_cell(w, c, max_chars, bold=True, fill="DDDDDD") for c in sh.columns])
            return w

        ws = new_ws(part)
        for rec in sh.data:
            if n >= MAX_ROWS:
                part += 1
                ws = new_ws(part)
                n = 0
            values = [rec.get(c) for c in sh.columns] if isinstance(rec, dict) else list(rec)
            fill = None
            if status_idx is not None:
                fill = STATUS_FILL.get(str(values[status_idx]).lower())
            ws.append([_cell(ws, v, max_chars, fill=fill if i == status_idx else None) for i, v in enumerate(values)])
            n += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sh.columns))}{max(1, n + 1)}"
    about = wb.create_sheet("About")
    for line in (f"Generated {now_iso()} by maintdoc {TOOL_VERSION}", DRAFT_DISCLAIMER,
                 "Full text and history are in maintenance.db (SQLite)."):
        about.append([_cell(about, line, max_chars)])
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(str(tmp))
    tmp.replace(path)
    return path


def _q(conn: sqlite3.Connection, sql: str, params=()) -> Iterable[dict]:
    cur = conn.execute(sql, params)
    cols = [d[0] for d in cur.description]
    for r in cur:
        yield dict(zip(cols, r))


def _cols(conn: sqlite3.Connection, table: str) -> list[str]:
    return [r[1] for r in conn.execute(f"PRAGMA table_info({table})")]


def export_source_register(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    src_cols = _cols(conn, "sources")
    return write_workbook(path, [
        Sheet("Sources", src_cols, _q(conn, "SELECT * FROM sources ORDER BY source_id"),
              {"rel_path": 60, "abs_path": 60, "metadata_origin_json": 80}, status_col="status"),
        Sheet("Versions", _cols(conn, "source_versions"),
              _q(conn, "SELECT * FROM source_versions ORDER BY source_id, version_id")),
        Sheet("Pages", [c for c in _cols(conn, "pages") if c != "page_text"],
              _q(conn, f"SELECT {', '.join(c for c in _cols(conn, 'pages') if c != 'page_text')} FROM pages "
                       f"ORDER BY source_id, page_no"), status_col="extraction_status"),
        Sheet("Status_Summary", ["status", "extraction_status", "count"],
              _q(conn, "SELECT status, extraction_status, COUNT(*) count FROM sources GROUP BY 1, 2")),
    ], int(cfg.get("export.max_cell_chars", 32000)))


def export_evidence_register(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    ev_cols = ["evidence_id", "source_id", "filename", "page_no", "seq", "block_type", "section_path", "text",
               "display_text", "chapter", "chapter_override", "classification_status", "chapter_rule", "equipment",
               "model", "component", "revision", "extraction_method", "ocr_confidence", "ocr_verified",
               "source_sha256", "status", "status_reason", "in_manual", "manual_exclusion", "dup_role",
               "canonical_evidence_id", "review_status", "reviewed_by", "reviewed_at", "review_comment", "safety_level",
               "quality_score", "bbox_x0", "bbox_y0", "bbox_x1", "bbox_y1", "bbox_approx", "related_evidence_id",
               "created_at"]
    sql = ("SELECT e.*, s.filename FROM evidence e JOIN sources s ON s.source_id=e.source_id "
           "ORDER BY e.source_id, e.page_no, e.seq")
    return write_workbook(path, [
        Sheet("Evidence", ev_cols, _q(conn, sql), {"text": 80, "display_text": 60, "section_path": 40,
                                                   "source_sha256": 66}, status_col="review_status"),
        Sheet("Quantities", _cols(conn, "quantities"), _q(conn, "SELECT * FROM quantities ORDER BY quantity_id")),
        Sheet("Exact_Duplicate_Groups", _cols(conn, "duplicate_groups"),
              _q(conn, "SELECT * FROM duplicate_groups ORDER BY group_id")),
        Sheet("Near_Duplicates", _cols(conn, "near_duplicates"),
              _q(conn, "SELECT * FROM near_duplicates ORDER BY score DESC")),
        Sheet("Review_Decisions", _cols(conn, "review_decisions"),
              _q(conn, "SELECT * FROM review_decisions ORDER BY decision_id")),
        Sheet("Visual_Checks", _cols(conn, "visual_checks"),
              _q(conn, "SELECT * FROM visual_checks ORDER BY source_id, page_no"), status_col="status"),
    ], int(cfg.get("export.max_cell_chars", 32000)))


def export_conflict_register(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    def detail():
        for c in rows(conn, "SELECT * FROM conflicts ORDER BY conflict_id"):
            for e in rows(conn, "SELECT e.evidence_id, e.source_id, s.filename, s.revision, e.page_no, e.model, "
                                "e.text, e.review_status FROM conflict_evidence ce JOIN evidence e ON "
                                "e.evidence_id=ce.evidence_id JOIN sources s ON s.source_id=e.source_id "
                                "WHERE ce.conflict_id=?", (c["conflict_id"],)):
                yield {"conflict_id": c["conflict_id"], "conflict_type": c["conflict_type"], "severity": c["severity"],
                       "status": c["status"], **e}
    return write_workbook(path, [
        Sheet("Conflicts", _cols(conn, "conflicts"), _q(conn, "SELECT * FROM conflicts ORDER BY blocking DESC, "
                                                              "severity, conflict_id"),
              {"description": 90, "values_json": 60, "evidence_ids": 50}, status_col="severity"),
        Sheet("Conflict_Evidence", ["conflict_id", "conflict_type", "severity", "status", "evidence_id", "source_id",
                                    "filename", "revision", "page_no", "model", "text", "review_status"], detail(),
              {"text": 90}),
    ], int(cfg.get("export.max_cell_chars", 32000)))


def export_error_register(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    codes = [{"error_code": k, "default_severity": v[0], "description": v[1]} for k, v in sorted(ERROR_CODES.items())]
    return write_workbook(path, [
        Sheet("Errors", _cols(conn, "extraction_errors"),
              _q(conn, "SELECT * FROM extraction_errors ORDER BY source_id, page_no, error_id"),
              {"message": 90, "details_json": 60}, status_col="severity"),
        Sheet("Summary", ["error_code", "severity", "status", "count"],
              _q(conn, "SELECT error_code, severity, status, COUNT(*) count FROM extraction_errors GROUP BY 1,2,3 "
                       "ORDER BY 1")),
        Sheet("Error_Codes", ["error_code", "default_severity", "description"], codes, {"description": 90}),
    ], int(cfg.get("export.max_cell_chars", 32000)))


def export_validation_report(report, cfg: Config, path: Path) -> Path:
    checks = [{"group": c.group, "check": c.name, "target": c.target, "status": c.status, "detail": c.detail}
              for c in report.checks]
    summary = [{"item": "Overall result", "value": report.overall},
               {"item": "Verification run", "value": report.run_id},
               {"item": "Generated at", "value": report.generated_at}] + \
              [{"item": f"Checks {k}", "value": v} for k, v in sorted(report.counts().items())]
    groups = sorted({c.group for c in report.checks})
    sheets = [Sheet("Acceptance_Criteria", ["no", "criterion", "status", "detail"], report.acceptance,
                    {"criterion": 55, "detail": 90}, status_col="status"),
              Sheet("Summary", ["item", "value"], summary, {"item": 30, "value": 60})]
    for g in groups:
        sheets.append(Sheet(g.capitalize(), ["group", "check", "target", "status", "detail"],
                            [c for c in checks if c["group"] == g], {"target": 35, "detail": 100},
                            status_col="status"))
    return write_workbook(path, sheets, int(cfg.get("export.max_cell_chars", 32000)))
