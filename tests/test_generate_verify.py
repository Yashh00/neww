from __future__ import annotations

import json

import pymupdf
import pytest
from docx import Document

from maintdoc.db import rows
from maintdoc.evidence import find_evidence_ids
from maintdoc.generate.manual import GenerationBlocked, build_manual
from maintdoc.generate.run import output_paths, run_generate
from maintdoc.inventory import run_inventory
from maintdoc.reporting.export import FILES, run_export
from maintdoc.testing.corpus import _manual
from maintdoc.verify import run_verify

from conftest import complete_review, pass_output_checks


def pdf_text(path):
    with pymupdf.open(str(path)) as d:
        return "\n".join(p.get_text() for p in d)


def test_draft_manual_every_item_cited(ws):
    res = run_generate(ws.conn, ws.cfg, "RUN-g", ["draft"])["draft"]
    assert res["status"] == "generated" and not res["errors"]
    paths = output_paths(ws.cfg, "draft")
    man = json.loads(paths["manifest"].read_text(encoding="utf-8"))
    assert man["items"] and all(it["citations"] for it in man["items"])
    text = pdf_text(paths["pdf"])
    for it in man["items"]:
        for cid in it["citations"]:
            assert cid in text.replace("\n", " ") or cid in text
    docx_text = "\n".join(p.text for p in Document(str(paths["docx"])).paragraphs)
    assert find_evidence_ids(docx_text)
    # different revision values are both present, each with its own citation
    assert "45 Nm (33 ft-lb)" in text and "50 Nm (37 ft-lb)" in text
    assert "DRAFT - NOT RELEASED" in text


def test_draft_manual_includes_exact_duplicate_citations(ws):
    manual = build_manual(ws.conn, ws.cfg, "draft")
    step1 = [i for i in manual.items() if i.text.startswith("1. Switch off the press")]
    assert len(step1) == 1  # shown once ...
    cites = {c.evidence_id for c in step1[0].citations}
    assert len(cites) >= 2  # ... citing every source revision that contains it
    roles = {c.role for c in step1[0].citations}
    assert "identical_file" in roles  # the byte-identical copy of Rev B is cited too


def test_approved_manual_blocked_by_open_critical_conflicts(ws):
    with pytest.raises(GenerationBlocked):
        build_manual(ws.conn, ws.cfg, "approved")
    res = run_generate(ws.conn, ws.cfg, "RUN-g", ["approved"])["approved"]
    assert res["status"] == "blocked"


def test_full_review_cycle_passes_verification(ws):
    complete_review(ws)
    res = run_generate(ws.conn, ws.cfg, "RUN-g")
    assert res["approved"]["status"] == "generated" and res["approved"]["stats"]["items"] > 0
    pass_output_checks(ws)
    rep = run_verify(ws.conn, ws.cfg, "RUN-v")
    statuses = {a["no"]: a["status"] for a in rep.acceptance}
    assert all(statuses[str(i)] == "pass" for i in range(1, 10)), rep.acceptance
    assert statuses["10"] == "pending_human"  # formal release is never automatic
    assert rep.overall.startswith("PASS") and "NOT RELEASED" in rep.overall
    approved_pdf = pdf_text(output_paths(ws.cfg, "approved")["pdf"])
    assert "50 Nm (37 ft-lb)" in approved_pdf and "45 Nm (33 ft-lb)" not in approved_pdf
    assert "APPROVED CONTENT - NOT RELEASED" in approved_pdf


def test_exports_written(ws):
    rep = run_verify(ws.conn, ws.cfg, "RUN-v")
    out = run_export(ws.conn, ws.cfg, "RUN-x", report=rep)
    for key, name in FILES.items():
        assert key in out and (ws.cfg.output_dir / name).exists(), name
    import openpyxl
    wb = openpyxl.load_workbook(ws.cfg.output_dir / "Validation_Report.xlsx", read_only=True)
    assert "Acceptance_Criteria" in wb.sheetnames
    ev = openpyxl.load_workbook(ws.cfg.output_dir / "Evidence_Register.xlsx", read_only=True)["Evidence"]
    header = [c.value for c in next(ev.iter_rows(max_row=1))]
    for col in ("evidence_id", "source_id", "page_no", "section_path", "text", "source_sha256", "extraction_method",
                "review_status"):
        assert col in header


def test_excel_text_is_never_a_formula(tmp_path):
    import openpyxl
    from maintdoc.reporting.xlsx import Sheet, write_workbook
    p = write_workbook(tmp_path / "x.xlsx", [Sheet("S", ["text"], [{"text": "=HYPERLINK(\"http://x\")"},
                                                                   {"text": "ok\x07bell"}])])
    ws = openpyxl.load_workbook(p)["S"]
    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("=HYPERLINK")
    assert ws["A3"].value == "okbell"


def test_source_change_invalidates_approvals_and_outputs(own_sources):
    w = own_sources
    complete_review(w)
    run_generate(w.conn, w.cfg, "RUN-g")
    approved_before = w.one("SELECT COUNT(*) FROM evidence WHERE review_status='approved'")
    assert approved_before > 0
    rev_c = w.corpus["rev_c"]
    _manual(rev_c, "C", "2023-06-01", "55", "41", "1000")  # engineering change in the source PDF
    # verification detects the change even before inventory runs
    rep = run_verify(w.conn, w.cfg, "RUN-v1")
    crit6 = [a for a in rep.acceptance if a["no"] == "6"][0]
    assert crit6["status"] == "fail"
    stats = run_inventory(w.conn, w.cfg, "RUN-i2", show_progress=False)
    assert stats.modified == 1 and stats.invalidated_approvals > 0
    sid = w.one("SELECT source_id FROM sources WHERE rel_path LIKE '%RevC.pdf'")
    assert w.one("SELECT COUNT(*) FROM evidence WHERE source_id=? AND review_status='approved'", (sid,)) == 0
    decisions = w.one("SELECT COUNT(*) FROM review_decisions WHERE decision='invalidate' AND reviewer='system'")
    assert decisions >= stats.invalidated_approvals
    # re-process: new conflict (45 vs 55 Nm) blocks the approved manual; stale outputs are marked outdated
    from maintdoc.analysis.validate import run_validate
    from maintdoc.extraction.runner import run_extraction
    run_extraction(w.conn, w.cfg, "RUN-e2", workers=1, show_progress=False)
    run_validate(w.conn, w.cfg, "RUN-v2", show_progress=False)
    res = run_generate(w.conn, w.cfg, "RUN-g2")
    assert res["approved"]["status"] == "blocked"
    assert not output_paths(w.cfg, "approved")["pdf"].exists()
    assert list(w.cfg.output_dir.glob("*APPROVED_OUTDATED.pdf"))
    torque = rows(w.conn, "SELECT * FROM conflicts WHERE conflict_type='torque' AND status='open'")
    assert torque and "55 Nm" in torque[0]["description"]
